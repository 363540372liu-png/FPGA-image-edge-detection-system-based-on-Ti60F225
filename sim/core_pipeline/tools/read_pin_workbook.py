import json
import re
import sys

import openpyxl


workbook = openpyxl.load_workbook(sys.argv[1], data_only=False, read_only=True)
print("SHEETS", workbook.sheetnames)
pattern = re.compile(r"(cmos|camera|sensor|dvp|pclk|href|vsync|data\[|\bD[0-7]\b)", re.I)

for sheet in workbook.worksheets:
    hits = []
    for row in sheet.iter_rows():
        values = [cell.value for cell in row]
        if pattern.search(" | ".join("" if value is None else str(value) for value in values)):
            hits.append((row[0].row, values))
    if hits:
        print("SHEET", sheet.title, "HITS", len(hits))
        for row_number, values in hits[:100]:
            print(row_number, json.dumps(values, ensure_ascii=False, default=str))
